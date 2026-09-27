"""Create synthetic viewer fixtures only; run with the workstation's conda y Python."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

root = Path(__file__).resolve().parents[1] / "target" / "verification-fixtures"
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("Mochi Native Verification", 0)
doc.add_paragraph("This document is synthetic test data. No user document is used.")
doc.add_paragraph("中文段落：验证字体与只读预览。")
doc.save(root / "sample.v1.docx")

book = Workbook()
sheet = book.active
sheet.title = "Overview"
sheet.append(["课程", "计划", "完成度"])
sheet.append(["操作系统", "2026-09-05", "80%"])
sheet.append(["算法设计", "2026-09-06", "45%"])
second = book.create_sheet("Details")
second.append(["Item", "Value"])
for i in range(1, 6001):
    second.append([f"Row {i}", i])
book.save(root / "sample.xlsx")

slides = Presentation()
for title, body in [("Mochi Native", "Synthetic PowerPoint preview fixture"), ("Second slide", "Verify multi-page preview and rendering")]:
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    shape.text_frame.text = title
    shape.text_frame.add_paragraph().text = body
slides.save(root / "sample.v1.pptx")

pdf = canvas.Canvas(str(root / "sample.pdf"), pagesize=(595, 842))
for i in range(1, 9):
    pdf.setFont("Helvetica", 24)
    pdf.drawString(48, 760, f"Mochi PDF verification - page {i}")
    pdf.rect(48, 480, 400, 200)
    pdf.setFont("Helvetica", 14)
    pdf.drawString(48, 450, "Drag annotations here, switch pages, then reopen.")
    pdf.showPage()
pdf.save()
print(root)
